"""생성 결과(out/*_v2.json)를 보기 쉬운 엑셀 한 파일로 내보낸다.

사용: python export_xlsx.py <out 폴더> <엑셀 파일 경로>
시트: 읽는 법 · 자산 · 라인 구성 · 서버·설치 SW · 통신 흐름 · 정답지(판정) · 수기 대장 · 오류 정답
모든 값은 가상 데이터다. 공지(CVE)는 '정답지' 시트에만 있다.
"""
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from render_report import stats

FONT = "Malgun Gothic"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
STATUS_FILL = {
    "affected_no_fix": "F4B6B6", "affected_scenario": "F8CBAD", "affected_conditional": "FFE699",
    "unknown": "D9D9D9", "fixed": "C6E0B4", "not_listed": "EDEDED",
}
COMPANIES = ("VMA", "VMB", "VMC")


def load(out, code, kind):
    return json.loads((out / f"{code.lower()}_{kind}.json").read_text(encoding="utf-8"))


def join(items):
    return ", ".join(str(x) for x in items) if items else ""


def write_sheet(wb, title, headers, rows, widths=None, fills=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    body = Font(name=FONT, size=10)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
            c.alignment = Alignment(vertical="top")
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        if w is None:
            longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if r[i - 1] is not None])
            w = min(max(longest * 1.6 + 2, 8), 46)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    if fills:
        col, mapping = fills
        idx = headers.index(col) + 1
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for c in row:
                if c.value in mapping:
                    c.fill = PatternFill("solid", fgColor=mapping[c.value])
    return ws


def asset_row(code, a):
    patch = a.get("patchLevel") or {}
    sup = a.get("supportEnd") or {}
    return [code, a["id"], a["siteId"], a["line"], a.get("zone"), a["name"], a["assetType"], a["category"],
            a["vendor"], a["product"], a.get("version"), a.get("firmwareVersion") or a.get("osBuild") or patch.get("label"),
            a.get("installedYear"), a.get("lineWave"), sup.get("date"), (a.get("discontinuedOn") or {}).get("date"),
            "종료" if a.get("supportStatus") == "ended" else "", a["criticality"], join(a.get("patchConstraints")),
            "가능" if a.get("rollbackPossible") else "불가", a.get("exposure"), a.get("edr"), a.get("rebootWindow"),
            a.get("ownerId"), a.get("ip"), a.get("networkSegmentId"), a.get("lifecycle"), a.get("lastVerifiedAt"),
            a.get("versionEvidence"), a.get("registerStatus"), a.get("process"), "예" if a.get("bottleneck") else "", a.get("redundancy"), a.get("osManagement")]


ASSET_HEAD = ["회사", "자산 ID", "사이트", "라인", "구역", "이름", "종류", "구분(IT/OT)", "제조사", "제품", "버전(기본)",
              "펌웨어·빌드(배정)", "도입 연도", "도입 시기", "지원 종료일", "단종일", "지원 상태", "중요도", "작업 제약",
              "되돌리기", "노출", "EDR", "재부팅 창", "담당", "IP", "망 구역", "수명주기", "마지막 확인", "버전 근거", "대장 상태", "공정", "병목", "이중화", "OS·펌웨어 관리"]


def main(out_dir, xlsx):
    out = pathlib.Path(out_dir)
    data = {c: load(out, c, "factory_v2") for c in COMPANIES}
    exp = {c: load(out, c, "expected_matches") for c in COMPANIES}
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "읽는 법"

    # ---- 읽는 법
    lines = [
        ("가상 공장 데이터셋 v2", True),
        ("회사·공장·자산은 모두 가상이다. 제품·버전 번호는 실제 공개 번호를 가정해 배정한 값이며 실제 회사·장비의 값이 아니다.", False),
        ("자산 시트에는 공지·CVE가 없다. 공지와의 비교 결과는 '정답지(판정)' 시트에만 있고, 수집기 결과를 비교하는 QA용이다. EPSS는 시험용 값이다.", False),
        ("", False),
        ("시트 안내", True),
        ("자산: 장비·서버 목록(제품, 펌웨어·버전, 도입 연도, 지원 종료일, 제약, 담당, IP).", False),
        ("라인 구성: 라인(공정)별로 어떤 제조사·OS를 쓰는지 한눈에.", False),
        ("서버·설치 SW: 서버의 OS와 응용 프로그램(OPC UA, SCADA, DB).", False),
        ("통신 흐름: 장비·서버 사이의 통신(포트는 예시).", False),
        ("정답지(판정): 제품·버전을 공개 공지와 비교한 기대 결과. 색으로 구분.", False),
        ("수기 대장: 엑셀·협력사 파일에서 모은 지저분한 원본 기록. 오류 정답 시트와 짝.", False),
        ("오류 정답: 수기 대장의 어떤 칸이 왜 틀렸고 맞는 값은 무엇인지.", False),
        ("", False),
        ("판정 의미", True),
        ("affected_scenario: 그 버전이면 영향 범위, 수정 버전 있음 / affected_conditional: 기능 전제가 맞을 때만 해당 / affected_no_fix: 영향 범위인데 수정 계획 없음", False),
        ("fixed: 수정 버전 이상 / not_listed: 공지의 영향 목록에 없음 / unknown: 영향 범위이나 수정 버전을 확인 못 함(확인 필요)", False),
        ("", False),
        ("회사별 수치", True),
    ]
    for text, bold in lines:
        ws0.append([text])
        ws0.cell(row=ws0.max_row, column=1).font = Font(name=FONT, bold=bold, size=14 if bold and text.startswith("가상") else 11)
    rows = {c: stats(data[c], exp[c]) for c in COMPANIES}
    ws0.append(["항목"] + list(COMPANIES))
    for c in ws0[ws0.max_row]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    for k in rows["VMA"]:
        ws0.append([k] + [str(rows[c][k]) for c in COMPANIES])
        for c in ws0[ws0.max_row]:
            c.font = Font(name=FONT, size=10)
    ws0.column_dimensions["A"].width = 60
    for col in "BCD":
        ws0.column_dimensions[col].width = 34

    # ---- 자산
    assets = [asset_row(c, a) for c in COMPANIES for a in data[c]["canonical_assets"]]
    write_sheet(wb, "자산", ASSET_HEAD, assets, widths={"이름": 30, "작업 제약": 28})

    # ---- 라인 구성
    line_rows = []
    for c in COMPANIES:
        by_line = {}
        for a in data[c]["canonical_assets"]:
            by_line.setdefault(a["line"], []).append(a)
        for l in data[c]["lines"]:
            q = by_line.get(l["id"], [])
            g = lambda t: join(sorted({a["product"] for a in q if a["assetType"] == t}))
            line_rows.append([c, l["siteId"], l["id"], l.get("lineType"), join(l["operations"]), l["productFamily"], q[0].get("lineWave") if q else "", g("CNC 컨트롤러"), g("PLC"),
                              g("인버터"), g("로봇 컨트롤러"), g("산업용 스위치"), g("HMI PC"), g("비전 검사 PC"), len(q)])
    write_sheet(wb, "라인 구성", ["회사", "사이트", "라인", "라인 유형", "공정", "제품군", "도입 시기", "CNC", "PLC(세척·이송)", "인버터", "로봇", "스위치", "HMI OS", "비전 PC OS", "자산 수"], line_rows)

    # ---- 서버·설치 SW
    sw_rows = []
    for c in COMPANIES:
        am = {a["id"]: a for a in data[c]["canonical_assets"]}
        for s in data[c]["software_installations"]:
            if not s.get("role"):
                continue
            a = am[s["assetId"]]
            sup = s.get("supportEnd") or {}
            sw_rows.append([c, a["id"], a["name"], a["assetType"], a["zone"], s["role"], s["vendor"], s["product"], s.get("version"),
                            (s.get("patchLevel") or {}).get("label"), sup.get("date"), sup.get("confidence"), a.get("installedYear")])
    write_sheet(wb, "서버·설치 SW", ["회사", "서버 ID", "서버 이름", "서버 종류", "구역", "구분(os/app)", "제조사", "제품", "버전", "빌드·패치", "지원 종료일", "종료일 근거", "서버 도입 연도"], sw_rows, widths={"서버 이름": 30})

    # ---- 통신 흐름
    flow_rows = []
    for c in COMPANIES:
        am = {a["id"]: a for a in data[c]["canonical_assets"]}
        for f in data[c]["application_flows"]:
            s, t = am[f["sourceAssetId"]], am[f["targetAssetId"]]
            flow_rows.append([c, f["id"], s["name"], s["assetType"], s["vendor"], t["name"], t["assetType"], f["protocol"], f["port"], f["purpose"], f.get("note")])
    write_sheet(wb, "통신 흐름", ["회사", "흐름 ID", "출발", "출발 종류", "출발 제조사", "도착", "도착 종류", "프로토콜", "포트", "목적", "비고"], flow_rows, widths={"목적": 40, "비고": 40})

    # ---- 정답지(판정)
    ans = []
    for c in COMPANIES:
        am = {a["id"]: a for a in data[c]["canonical_assets"]}
        for item in exp[c]["items"]:
            a = am[item["assetId"]]
            if not item["matches"]:
                ans.append([c, a["id"], a["name"], item.get("softwareId"), item["product"], item["version"], "(retired — 판정 없음)", "", "", "", "", "", "", "", ""])
            for m in item["matches"]:
                cv = m.get("cvss") or {}
                ans.append([c, a["id"], a["name"], item.get("softwareId"), item["product"], item["version"], m["cve"], m["status"],
                            "예" if m["kev"] else "아니오", m.get("kevAdded"), cv.get("score"), (m.get("epss") or {}).get("score"),
                            m.get("confidence"), m.get("precondition"), m["url"]])
    write_sheet(wb, "정답지(판정)", ["회사", "자산 ID", "자산 이름", "설치 SW ID", "제품 키", "배정 버전", "CVE", "판정", "KEV 등재", "KEV 등재일",
                                  "CVSS", "EPSS(시험값)", "근거 신뢰도", "전제 조건", "공지 주소"], ans, widths={"자산 이름": 30, "공지 주소": 50},
                fills=("판정", STATUS_FILL))

    # ---- 수기 대장 / 오류 정답
    keys = ["설비명", "제조사", "제품", "버전", "IP", "담당", "라인", "시리얼", "상태", "확인일"]
    raw_rows, issue_rows = [], []
    for c in COMPANIES:
        r2a = data[c]["quality_truth"]["row_to_asset"]
        for r in data[c]["raw_asset_register"]:
            raw_rows.append([c, r["id"], r["source"], r["sourceRow"], r2a.get(r["id"])] + [r["payload"].get(k) for k in keys])
        for i in data[c]["quality_truth"]["issues"]:
            issue_rows.append([c, i["rowId"], i["field"], i["issueType"], i["expected"], i["reported"], i.get("resolution")])
    write_sheet(wb, "수기 대장", ["회사", "행 ID", "출처 파일", "행 번호", "정답 자산 ID"] + keys, raw_rows, widths={"설비명": 36})
    write_sheet(wb, "오류 정답", ["회사", "행 ID", "틀린 칸", "오류 유형", "맞는 값", "대장에 적힌 값", "처리"], issue_rows)

    wb.save(xlsx)
    print("saved", xlsx, {ws.title: ws.max_row - 1 for ws in wb.worksheets[1:]})


if __name__ == "__main__":
    main(*sys.argv[1:3])
